"""
Austin FC TVOF | Fan linkage workbook (v1, 2026-09-23)
-------------------------------------------------------------
Row-level workbook for tracing fans across the client tables. One sheet per
view; add new views as build_<name>() functions and register them in SHEETS.

Sheets:
    Dup_internal_fan_id   every Fan Info row whose internal_fan_id appears more
                        than once, grouped together, with all other columns
                        and the fields that differ inside each group
    Dup_fan_connections   the same rows, with what each one links to in
                        Activation, Attendance and Sales (direct and through
                        the seat bridge); missing links are shaded red
    Dup_fan_combined      one row per duplicate fan (group_no): all of its rows
                        rolled up, activity summed across its accounts
    Connection_Summary    how many duplicate fans / rows have each link
    Activation_Attendance one row per World Cup activation fan, followed through
                        Fan Info into Attendance: scans and matches BEFORE and
                        AFTER their first signup
    Activation_Summary    activation fans by attendance status, overall and by event

Usage:
    python tvof_fan_linkage.py --data-dir "<your folder of client CSVs>"

The Sales bridge columns need account_bridge_LOCAL_ONLY.csv (from
tvof_account_bridge.py) in the data folder; without it they stay empty.

Output: fan_linkage_LOCAL_ONLY.xlsx in the data folder.
ROW LEVEL (hashed IDs): keep local, do not share or commit.
"""
import argparse, functools, os, sys

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tvof_data_audit as A   # noqa: E402

FAN_DATES = ["seatgeek_since_date", "seatgeek_stm_since_date"]
BAND = PatternFill("solid", start_color="EAF1FB")
HEAD = PatternFill("solid", start_color="1F3864")
GROUP_EDGE = Border(top=Side(style="thin", color="808080"))
LINKED = PatternFill("solid", start_color="C6EFCE")
MISSING = PatternFill("solid", start_color="FFC7CE")
# count columns shaded green (>0) or red (0) on the connections sheet
LINK_COLS = {"activation_signups", "attendance_scans", "sales_rows_direct", "sales_rows_via_bridge"}
BRIDGE_FILE = "account_bridge_LOCAL_ONLY.csv"
ORIGINAL_SALE = {"ticket", "subscription"}   # revenue definition B in the audit


def read_fan(path):
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding_errors="replace")
    df = df.apply(lambda s: s.str.strip())
    df = df.mask(df.apply(lambda s: s.str.lower().isin(A.NULL_TOKENS)))
    for c in FAN_DATES:
        if c in df:
            df[c] = pd.to_datetime(df[c], format="%m/%d/%Y", errors="coerce").dt.date
    return df


def read_cols(path, cols, chunksize=500_000):
    """Selected columns as stripped strings, null tokens -> NaN, read in chunks."""
    parts = []
    for ch in pd.read_csv(path, usecols=cols, dtype=str, keep_default_na=False,
                        chunksize=chunksize, encoding_errors="replace"):
        ch = ch.apply(lambda s: s.str.strip())
        parts.append(ch.mask(ch.apply(lambda s: s.str.lower().isin(A.NULL_TOKENS))))
    return pd.concat(parts, ignore_index=True)


@functools.lru_cache(maxsize=None)
def dup_rows(fan_path):
    """Fan Info rows sharing an internal_fan_id, one block per id."""
    fan = read_fan(fan_path)
    key = "internal_fan_id"
    d = fan[fan[key].notna() & fan[key].duplicated(keep=False)].copy()
    d = d.sort_values([key, "internal_account_id", "mls_id"], na_position="last")
    g = d.groupby(key, sort=False)
    d.insert(0, "group_no", g.ngroup() + 1)
    d.insert(1, "rows_in_group", g[key].transform("size"))
    other = [c for c in fan.columns if c != key]
    differs = g[other].nunique(dropna=False)
    differs = differs.apply(lambda r: ", ".join(c for c in other if r[c] > 1) or "(identical rows)", axis=1)
    d.insert(2, "fields_that_differ", d[key].map(differs))
    cols = ["group_no", "rows_in_group", "fields_that_differ", key] + other
    A.log(f"  {key}: {d[key].nunique():,} duplicated ids across {len(d):,} rows")
    return d[cols]


def build_dup_fan_id(files, a):
    return dup_rows(files["fan"])


def sales_for(path, accounts):
    """Rows and original-sale total_payment per internal_account_id, for the given accounts only."""
    acc, it, pay = "internal_account_id", "item_type", "total_payment"
    parts = []
    for ch in pd.read_csv(path, usecols=[acc, it, pay], dtype=str, keep_default_na=False,
                        chunksize=500_000, encoding_errors="replace"):
        ch[acc] = ch[acc].str.strip()
        ch = ch[ch[acc].isin(accounts)]
        if len(ch):
            parts.append(ch)
    if not parts:
        return pd.DataFrame(columns=["rows", "orig_payment"])
    s = pd.concat(parts, ignore_index=True)
    s["orig"] = pd.to_numeric(s[pay], errors="coerce").where(s[it].str.strip().str.lower().isin(ORIGINAL_SALE))
    return s.groupby(acc).agg(rows=(acc, "size"), orig_payment=("orig", "sum"))


@functools.lru_cache(maxsize=None)
def connections(fan_path, act_path, att_path, sales_path, bridge_path):
    """One row per duplicate-fan row: what it links to in each table, and what it doesn't."""
    d = dup_rows(fan_path)[["group_no", "rows_in_group", "internal_fan_id", "internal_account_id", "mls_id"]].copy()
    fids, accs = set(d.internal_fan_id.dropna()), set(d.internal_account_id.dropna())

    A.log("  activation ...")
    act = read_cols(act_path, ["internal_fan_id", "ticket_scan_datetime"])
    act = act[act.internal_fan_id.isin(fids)]
    ag = act.groupby("internal_fan_id").agg(activation_signups=("internal_fan_id", "size"),
                                            activation_scanned=("ticket_scan_datetime", "count"))
    d = d.join(ag, on="internal_fan_id")

    A.log("  attendance ...")
    att = read_cols(att_path, ["internal_account_id", "EventKey", "SeasonKey"])
    att = att[att.internal_account_id.isin(accs)]
    tg = att.groupby("internal_account_id").agg(
        attendance_scans=("EventKey", "size"), attendance_matches=("EventKey", "nunique"),
        attendance_seasons=("SeasonKey", lambda s: ", ".join(sorted(s.dropna().unique()))))
    d = d.join(tg, on="internal_account_id")
    # distinct matches per fan across all of its accounts (two accounts can scan the same match)
    af = att.merge(d[["internal_account_id", "internal_fan_id"]].drop_duplicates(), on="internal_account_id")
    fan_matches = af.groupby("internal_fan_id").EventKey.nunique()

    if os.path.exists(bridge_path):
        br = pd.read_csv(bridge_path, dtype=str)
        br = br[br.fan_account_id.isin(accs)]
    else:
        A.log(f"  {BRIDGE_FILE} not found: bridge columns left empty (run tvof_account_bridge.py)")
        br = pd.DataFrame(columns=["sales_account_id", "fan_account_id", "events_together", "confidence"])
    br = br.drop_duplicates("fan_account_id").set_index("fan_account_id")
    d["bridge_sales_account"] = d.internal_account_id.map(br["sales_account_id"])
    d["bridge_confidence"] = d.internal_account_id.map(br["confidence"]).fillna("(no link)")
    d["bridge_events_together"] = pd.to_numeric(d.internal_account_id.map(br["events_together"]))
    # v2 bridge: usable = High/Medium on seats, or date-confirmed (first sale == seatgeek_since)
    if "usable" in br:
        d["bridge_usable"] = d.internal_account_id.map(br["usable"].astype(str).str.lower().eq("true")).fillna(False)
    else:
        d["bridge_usable"] = d.bridge_confidence.isin(["High", "Medium"])

    A.log("  sales (full file scan, 1 to 2 min) ...")
    sa = sales_for(sales_path, accs | set(d.bridge_sales_account.dropna()))
    d["sales_rows_direct"] = d.internal_account_id.map(sa["rows"])
    d["sales_rows_via_bridge"] = d.bridge_sales_account.map(sa["rows"])
    d["sales_orig_payment_via_bridge"] = pd.to_numeric(d.bridge_sales_account.map(sa["orig_payment"])).round(2)

    for c in ["activation_signups", "activation_scanned", "attendance_scans", "attendance_matches",
            "sales_rows_direct", "sales_rows_via_bridge"]:
        d[c] = d[c].fillna(0).astype(int)

    def status(r):
        has, miss = [], []
        (has if r.activation_signups else miss).append("Activation")
        (has if r.attendance_scans else miss).append("Attendance")
        if r.sales_rows_direct:
            has.append("Sales")
        elif r.sales_rows_via_bridge and r.bridge_usable:
            has.append(f"Sales (bridge, {r.bridge_confidence})")
        elif r.sales_rows_via_bridge:
            miss.append("Sales (bridge Low, not date-confirmed)")
        else:
            miss.append("Sales")
        return pd.Series({"linked_to": ", ".join(has) or "(nothing)", "not_linked_to": ", ".join(miss) or "-"})

    d = pd.concat([d, d.apply(status, axis=1)], axis=1)
    lead = ["group_no", "rows_in_group", "internal_fan_id", "internal_account_id", "linked_to", "not_linked_to"]
    return d[lead + [c for c in d.columns if c not in lead]], fan_matches


def _conn(files, a):
    bridge = os.path.join(a.data_dir, BRIDGE_FILE)
    return connections(files["fan"], files["activation"], files["attendance"], files["sales"], bridge)


def build_connections(files, a):
    return _conn(files, a)[0]


def _join_unique(s, sep=", "):
    vals = sorted({v.strip() for x in s.dropna() for v in str(x).split(",") if v.strip()})
    return sep.join(vals)


def build_combined(files, a):
    """One row per duplicate fan: its Fan Info rows and connections rolled up."""
    d, fan_matches = _conn(files, a)
    f = dup_rows(files["fan"])
    usable = d.bridge_usable
    d = d.assign(usable_rows=d.sales_rows_via_bridge.where(usable, 0),
                 usable_payment=d.sales_orig_payment_via_bridge.where(usable))
    g = d.groupby("group_no", sort=True)
    out = g.agg(internal_fan_id=("internal_fan_id", "first"),
                fan_info_rows=("rows_in_group", "first"),
                accounts=("internal_account_id", "nunique"),
                accounts_with_attendance=("attendance_scans", lambda s: int((s > 0).sum())),
                mls_ids=("mls_id", "nunique"),
                # activation is keyed on internal_fan_id, so every row of a fan repeats it: take it once
                activation_signups=("activation_signups", "first"),
                activation_scanned=("activation_scanned", "first"),
                attendance_scans=("attendance_scans", "sum"),
                attendance_seasons=("attendance_seasons", _join_unique),
                sales_rows_direct=("sales_rows_direct", "sum"),
                bridge_links=("bridge_confidence", lambda s: _join_unique(s[s.ne("(no link)")]) or "(no link)"),
                sales_rows_via_bridge=("usable_rows", "sum"),
                sales_orig_payment_via_bridge=("usable_payment", lambda s: round(s.sum(min_count=1), 2)))
    out.insert(out.columns.get_loc("attendance_seasons"), "attendance_matches",
               out.internal_fan_id.map(fan_matches).fillna(0).astype(int))
    fg = f.groupby("group_no")
    out["earliest_seatgeek_since"] = fg.seatgeek_since_date.min()
    out["seatgeek_zip_codes"] = fg.seatgeek_zip_code.agg(_join_unique)
    out["mls_zip_codes"] = fg.mls_zip_code.agg(_join_unique)
    out["marketing_optin"] = fg.mls_club_fan_marketing_optin_flag.agg(_join_unique)

    def status(r):
        has, miss = [], []
        (has if r.activation_signups else miss).append("Activation")
        (has if r.attendance_scans else miss).append("Attendance")
        if r.sales_rows_direct:
            has.append("Sales")
        elif r.sales_rows_via_bridge:
            has.append("Sales (bridge)")
        else:
            miss.append("Sales")
        return pd.Series({"linked_to": ", ".join(has) or "(nothing)", "not_linked_to": ", ".join(miss) or "-"})

    out = pd.concat([out, out.apply(status, axis=1)], axis=1).reset_index()
    lead = ["group_no", "internal_fan_id", "fan_info_rows", "accounts", "linked_to", "not_linked_to"]
    A.log(f"  {len(out):,} fans combined from {len(d):,} rows")
    return out[lead + [c for c in out.columns if c not in lead]]


STATUS_ORDER = ["Not in Fan Info (no SeatGeek link)", "In Fan Info, no ticketing account",
                "Has account, never attended", "Attended before signup only",
                "Attended before and after signup", "First attended AFTER signup"]


@functools.lru_cache(maxsize=None)
def activation_follow(fan_path, act_path, att_path):
    """Activation fan -> Fan Info.internal_fan_id -> internal_account_id -> Attendance, split at first signup."""
    act = read_cols(act_path, ["internal_fan_id", "signup_datetime", "event_name", "ticket_scan_datetime"])
    act = act[act.internal_fan_id.notna()].copy()
    act["signup"] = A._parse_dt(act.signup_datetime)
    names = lambda x: "; ".join(sorted(x.dropna().unique()))
    f = act.groupby("internal_fan_id").agg(
        activation_signups=("event_name", "size"), activation_events=("event_name", names),
        activation_scanned=("ticket_scan_datetime", "count"), first_signup=("signup", "min"))

    fan = read_fan(fan_path)[["internal_fan_id", "internal_account_id"]].dropna(subset=["internal_fan_id"])
    links = fan.dropna(subset=["internal_account_id"]).drop_duplicates()
    f["in_fan_info"] = f.index.isin(set(fan.internal_fan_id))
    f["seatgeek_accounts"] = links.groupby("internal_fan_id").size().reindex(f.index).fillna(0).astype(int)

    A.log("  attendance for activation fans ...")
    att = read_cols(att_path, ["internal_account_id", "EventKey", "AttendedDatetime"])
    att = att[att.internal_account_id.isin(set(links.internal_account_id))]
    m = att.merge(links, on="internal_account_id")
    m = m[m.internal_fan_id.isin(f.index)].copy()
    m["t"] = A._parse_dt(m.AttendedDatetime)
    m["before"] = m.t < m.internal_fan_id.map(f.first_signup)
    g = m.groupby("internal_fan_id")
    f["attendance_scans_before"] = m[m.before].groupby("internal_fan_id").size()
    f["attendance_scans_after"] = m[~m.before].groupby("internal_fan_id").size()
    f["matches_before"] = m[m.before].groupby("internal_fan_id").EventKey.nunique()
    f["matches_after"] = m[~m.before].groupby("internal_fan_id").EventKey.nunique()
    f["first_scan"] = g.t.min()
    f["last_scan"] = g.t.max()
    for c in ["attendance_scans_before", "attendance_scans_after", "matches_before", "matches_after"]:
        f[c] = f[c].fillna(0).astype(int)

    b, af = f.matches_before.gt(0), f.matches_after.gt(0)
    f["status"] = STATUS_ORDER[2]
    f.loc[b & ~af, "status"] = STATUS_ORDER[3]
    f.loc[b & af, "status"] = STATUS_ORDER[4]
    f.loc[~b & af, "status"] = STATUS_ORDER[5]
    f.loc[f.seatgeek_accounts.eq(0), "status"] = STATUS_ORDER[1]
    f.loc[~f.in_fan_info, "status"] = STATUS_ORDER[0]

    # matches played after the earliest signup: the ceiling for "after" attendance
    ev = att.assign(t=A._parse_dt(att.AttendedDatetime)).groupby("EventKey").t.min()
    later = int((ev > f.first_signup.min()).sum())
    for c in ["first_signup", "first_scan", "last_scan"]:
        f[c] = f[c].dt.tz_localize(None) if getattr(f[c].dt, "tz", None) else f[c]
    A.log(f"  {len(f):,} activation fans; {later} matches in Attendance after the first signup")
    return f.reset_index(), later, act[["internal_fan_id", "event_name"]].drop_duplicates()


def _act(files):
    return activation_follow(files["fan"], files["activation"], files["attendance"])


def build_activation(files, a):
    f = _act(files)[0]
    lead = ["internal_fan_id", "status", "in_fan_info", "seatgeek_accounts"]
    return f[lead + [c for c in f.columns if c not in lead]].sort_values(
        ["status", "matches_after"], ascending=[True, False],
        key=lambda s: s.map(STATUS_ORDER.index) if s.name == "status" else s)


def build_activation_summary(files, a):
    f, later, fe = _act(files)
    rows = []
    for st in STATUS_ORDER:
        m = f.status.eq(st)
        rows.append({"event": "ALL activation fans", "status": st, "fans": int(m.sum()),
                     "pct_of_fans": round(m.mean(), 4), "matches_after_total": int(f.matches_after[m].sum())})
    ef = fe.merge(f[["internal_fan_id", "status", "matches_after"]], on="internal_fan_id")
    for (e, st), x in ef.groupby(["event_name", "status"]):
        rows.append({"event": e, "status": st, "fans": len(x),
                     "pct_of_fans": round(len(x) / ef.event_name.eq(e).sum(), 4),
                     "matches_after_total": int(x.matches_after.sum())})
    out = pd.DataFrame(rows)
    out["_o"] = out.status.map(STATUS_ORDER.index)
    out["_e"] = out.event.ne("ALL activation fans")
    out = out.sort_values(["_e", "event", "_o"]).drop(columns=["_o", "_e"])
    notes = pd.DataFrame({"event": [
        "Route: Activation.internal_fan_id -> Fan Info.internal_fan_id -> internal_account_id (all of a fan's accounts) -> Attendance.",
        "Before / after = scan time vs the fan's FIRST activation signup_datetime.",
        f"Only {later} matches in Attendance were played after the earliest signup, so 'after' covers a short window.",
        "Attendance is one row per scanned ticket, so one fan can have several scans per match; matches_* counts distinct matches.",
        "A fan can sign up for several events, so the per-event rows add up to more than the fan total."]})
    return pd.concat([out, notes], ignore_index=True)


def build_summary(files, a):
    d = _conn(files, a)[0]
    hm = d.bridge_usable & d.sales_rows_via_bridge.gt(0)
    checks = [
        ("Activation (by internal_fan_id)", d.activation_signups.gt(0)),
        ("Attendance (by internal_account_id)", d.attendance_scans.gt(0)),
        ("Sales direct (by internal_account_id)", d.sales_rows_direct.gt(0)),
        ("Sales via seat bridge, usable (High/Medium or date-confirmed)", hm),
        ("Sales via seat bridge, any confidence", d.sales_rows_via_bridge.gt(0)),
        ("Sales by any usable route (direct or usable bridge)", d.sales_rows_direct.gt(0) | hm),
        ("Linked to nothing", d.linked_to.eq("(nothing)")),
    ]
    rows = []
    for name, m in checks:
        by_fan = m.groupby(d.group_no)
        fans = by_fan.all() if name == "Linked to nothing" else by_fan.any()
        rows.append({"connection": name, "fan_rows": int(m.sum()), "of_rows": len(d),
                    "pct_rows": round(m.mean(), 4), "duplicate_fans": int(fans.sum()),
                    "of_fans": len(fans), "pct_fans": round(fans.mean(), 4)})
    rows.append({"connection": "A fan counts as linked if ANY of its rows links (for 'Linked to nothing': if ALL rows link to nothing). "
                            "Bridge = account_bridge_LOCAL_ONLY.csv as last built by tvof_account_bridge.py."})
    return pd.DataFrame(rows)


# sheet name -> builder(files, args) -> DataFrame (group_no column optional)
SHEETS = {
    "Dup_internal_fan_id": build_dup_fan_id,
    "Dup_fan_connections": build_connections,
    "Dup_fan_combined": build_combined,
    "Connection_Summary": build_summary,
    "Activation_Attendance": build_activation,
    "Activation_Summary": build_activation_summary,
}


def style(ws, df):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
    if "group_no" in df:
        grp = df["group_no"].to_numpy()
        for i, gno in enumerate(grp, start=2):
            new = i == 2 or gno != grp[i - 3]
            for c in ws[i]:
                if gno % 2 == 0:
                    c.fill = BAND
                if new:
                    c.border = GROUP_EDGE
    for j, col in enumerate(df.columns, start=1):
        if col in LINK_COLS:
            for i in range(2, len(df) + 2):
                c = ws.cell(i, j)
                c.fill = LINKED if c.value else MISSING
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(70, max(10, width))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default=A.DEFAULT_DIR)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    files = A.find_files(a.data_dir)
    for k in ("fan", "activation", "attendance", "sales"):
        if k not in files:
            sys.exit(f"{k} file not found in {a.data_dir}")
    out = a.out or os.path.join(a.data_dir, "fan_linkage_LOCAL_ONLY.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, build in SHEETS.items():
            A.log(f"Building {name} ...")
            df = build(files, a)
            df.to_excel(xw, sheet_name=name[:31], index=False)
            style(xw.sheets[name[:31]], df)
    A.autofilter(out)
    A.log(f"Wrote {out}  (ROW-LEVEL: keep local)")


if __name__ == "__main__":
    main()
